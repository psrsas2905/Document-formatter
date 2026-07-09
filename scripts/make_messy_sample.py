"""Generate samples/input_messy.docx — the deterministic test fixture.

Reproduces the bad habits the pipeline must untangle:
  - headings faked with hand-applied bold + font size (no named styles)
  - a hand-bolded heading at body size (only "short + bold + no period" gives it away)
  - tab characters used for indentation
  - blank paragraphs from double-Enter used as spacing
  - bullets typed with mixed glyphs (bullet, dash, asterisk) and manual 1. / a) numbering
  - "Figure N" / "Table N" caption lines
  - one paragraph that already carries a real named style (must be trusted as-is)
  - an indented italic quote

Run from the repo root:  python scripts/make_messy_sample.py
"""

from __future__ import annotations

from pathlib import Path

import docx
from docx.shared import Pt


def para(doc, text: str, size: float | None = None, bold: bool = False, italic: bool = False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    if size is not None:
        run.font.size = Pt(size)
    run.font.bold = bold or None
    run.font.italic = italic or None
    return p


def main() -> None:
    out = Path("samples/input_messy.docx")
    out.parent.mkdir(parents=True, exist_ok=True)
    doc = docx.Document()

    # Fake title: biggest hand-set size, all caps, bold.
    para(doc, "ACME WIDGET INSTALLATION GUIDE", size=18, bold=True)
    para(doc, "")  # double-Enter spacing
    para(doc, "")

    # Fake H2: mid-size bold short line.
    para(doc, "Overview", size=14, bold=True)
    para(
        doc,
        "\tThe Acme Widget is a modular fastening system for industrial shelving. "
        "This guide walks a technician through unpacking, mounting and first-use "
        "checks. Read the safety section before starting work.",
    )
    para(doc, "")

    # Hand-bolded heading at BODY size — only short+bold+no-period identifies it.
    para(doc, "Safety Precautions", bold=True)
    para(
        doc,
        "Always disconnect mains power before opening the housing. The widget "
        "carries residual charge for up to two minutes after shutdown.",
    )
    # Mixed bullet glyphs, typed by hand.
    para(doc, "• Wear insulated gloves rated to 1000 V")
    para(doc, "- Keep the work area dry at all times")
    para(doc, "* Never bypass the interlock switch")
    para(doc, "")
    para(doc, "")

    para(doc, "Installation Steps", size=14, bold=True)
    para(doc, "1. Unpack the widget and check the contents against the packing list.")
    para(doc, "2. Mount the base plate using the four M6 bolts supplied.")
    para(doc, "\ta) Torque each bolt to the value shown below.")
    para(doc, "\tb) Re-check torque after the first 24 hours of operation.")
    para(doc, "3. Connect the signal cable to port A.")
    para(doc, "")

    # Captions.
    para(doc, "Figure 1: Wiring diagram for single-phase supply", size=9)
    para(
        doc,
        "The diagram above shows the recommended earthing arrangement. Local "
        "regulations take precedence where they differ.",
    )
    para(doc, "Table 2 - Torque settings by bolt size", size=9)
    para(doc, "")

    # A paragraph that already carries a real named style — must be trusted.
    doc.add_paragraph("Maintenance Schedule", style="Heading 2")
    para(
        doc,
        "Inspect the fastening ring every six months. Replace the gasket kit "
        "whenever the housing is opened.",
    )
    para(doc, "")

    # Indented italic quote.
    q = para(
        doc,
        "\tWarranty is void if the unit is operated outside the rated "
        "temperature envelope.",
        italic=True,
    )
    q.paragraph_format.left_indent = Pt(36)

    para(
        doc,
        "For spare parts, quote the serial number printed on the base plate when "
        "contacting Acme support.",
    )

    doc.save(out)
    print(f"Wrote {out} ({len(doc.paragraphs)} paragraphs)")


if __name__ == "__main__":
    main()
