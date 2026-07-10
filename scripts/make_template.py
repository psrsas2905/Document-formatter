"""Generate templates/org_standard.docx — a BRANDED stand-in organization template.

Real workflow: each organization supplies its own .docx/.dotx template (their
brand format, logo included) and the profile's template_file — or the CLI's
--template-docx flag — points at it. docformat pours content into that
template's named styles and never touches its headers/footers, so the brand
logo survives into every output.

Until a real org template is supplied, this script builds a demo one:
  - a generated logo image + company name in the page header
  - brand-colored Heading 1/2/3 styles
  - all styles the example profile maps to

Run from the repo root:  python scripts/make_template.py
"""

from __future__ import annotations

import struct
import zlib
from pathlib import Path

import docx
from docx.shared import Pt, RGBColor

REQUIRED_STYLES = [
    "Heading 1",
    "Heading 2",
    "Heading 3",
    "Body Text",
    "Caption",
    "List Bullet",
    "Quote",
]

BRAND_RGB = (0x1F, 0x4E, 0x79)  # Acme dark blue
ACCENT_RGB = (0xC0, 0x50, 0x2D)  # Acme rust


def write_logo_png(path: Path, w: int = 96, h: int = 96) -> None:
    """Write a simple two-tone geometric logo PNG without any imaging library."""
    rows = []
    for y in range(h):
        row = bytearray([0])  # filter byte: None
        for x in range(w):
            on_diagonal = abs((x * h // w) - y) < h // 6
            row += bytes(ACCENT_RGB if on_diagonal else BRAND_RGB)
        rows.append(bytes(row))

    def chunk(tag: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data))
            + tag
            + data
            + struct.pack(">I", zlib.crc32(tag + data))
        )

    ihdr = struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)  # 8-bit RGB
    png = (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", zlib.compress(b"".join(rows)))
        + chunk(b"IEND", b"")
    )
    path.write_bytes(png)


def main() -> None:
    out = Path("templates/org_standard.docx")
    logo = Path("templates/assets/acme_logo.png")
    logo.parent.mkdir(parents=True, exist_ok=True)
    out.parent.mkdir(parents=True, exist_ok=True)
    write_logo_png(logo)

    doc = docx.Document()
    available = {s.name for s in doc.styles}
    missing = [s for s in REQUIRED_STYLES if s not in available]
    if missing:
        raise SystemExit(f"Default template is missing required styles: {missing}")

    # Brand the heading styles (this is the template designer's job, done once
    # per organization — docformat itself never invents formatting).
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].font.color.rgb = RGBColor(*BRAND_RGB)

    # Brand header: logo + company wordmark. Lives in the template's default
    # header, which docformat preserves verbatim.
    header = doc.sections[0].header
    para = header.paragraphs[0]
    para.add_run().add_picture(str(logo), width=Pt(24), height=Pt(24))
    wordmark = para.add_run("  ACME CORP")
    wordmark.font.bold = True
    wordmark.font.color.rgb = RGBColor(*BRAND_RGB)

    doc.save(out)
    print(f"Wrote {out} (branded header logo: {logo})")


if __name__ == "__main__":
    main()
