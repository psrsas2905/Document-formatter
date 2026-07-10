"""Content-fidelity regressions from the expert review: tracked changes,
fields, content controls, native lists, hyperlinks, comments, endnotes and
table relationship leaks must never be lost SILENTLY."""

import shutil
import zipfile

import docx
import pytest
from docx.oxml import parse_xml
from docx.oxml.ns import qn

from docformat.apply import apply_styles
from docformat.classify import CONFIDENCE_THRESHOLD, classify
from docformat.ingest import ingest
from docformat.models import Block, BlockType, Document
from docformat.template import load_profile

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
R = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
RT_HYPERLINK = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"
)


@pytest.fixture()
def profile():
    return load_profile("config/template_profile.example.yaml")


def _pipeline(src_path, profile, tmp_path):
    doc = classify(ingest(src_path))
    styled = apply_styles(doc, profile, tmp_path / "out.docx")
    return doc, docx.Document(str(styled)), styled


def _inject_part(path, partname, xml, content_type, reltype):
    """Add a part + relationship + content-type override to a saved docx."""
    tmp = path.with_suffix(".tmp.docx")
    rid = "rIdInjected1"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w") as zout:
        for item in zin.namelist():
            data = zin.read(item)
            if item == "[Content_Types].xml":
                data = data.replace(
                    b"</Types>",
                    f'<Override PartName="/word/{partname}" ContentType="{content_type}"/>'
                    "</Types>".encode(),
                )
            elif item == "word/_rels/document.xml.rels":
                data = data.replace(
                    b"</Relationships>",
                    f'<Relationship Id="{rid}" Type="{reltype}" Target="{partname}"/>'
                    "</Relationships>".encode(),
                )
            zout.writestr(item, data)
        zout.writestr(f"word/{partname}", xml)
    shutil.move(tmp, path)


def test_tracked_insertion_kept_deletion_accepted(tmp_path, profile):
    src = docx.Document()
    p = src.add_paragraph("Before ")
    p._p.append(parse_xml(
        f'<w:ins {W} w:id="1" w:author="Rev" w:date="2026-01-01T00:00:00Z">'
        "<w:r><w:t>INSERTED</w:t></w:r></w:ins>"
    ))
    p._p.append(parse_xml(
        f'<w:del {W} w:id="2" w:author="Rev" w:date="2026-01-01T00:00:00Z">'
        "<w:r><w:delText>DELETED</w:delText></w:r></w:del>"
    ))
    p.add_run(" after.")
    src_path = tmp_path / "src.docx"
    src.save(src_path)

    doc, out, _ = _pipeline(src_path, profile, tmp_path)
    text = "\n".join(p.text for p in out.paragraphs)
    assert "Before INSERTED after." in text
    assert "DELETED" not in text
    assert any("Tracked changes" in n for n in doc.notes)


def test_field_frozen_to_cached_text(tmp_path, profile):
    src = docx.Document()
    p = src.add_paragraph("As shown in ")
    p._p.append(parse_xml(
        f'<w:fldSimple {W} w:instr=" REF _Ref123 \\h "><w:r><w:t>Section 2.1</w:t></w:r></w:fldSimple>'
    ))
    p.add_run(" above, values hold.")
    src_path = tmp_path / "src.docx"
    src.save(src_path)

    doc, out, _ = _pipeline(src_path, profile, tmp_path)
    assert any("As shown in Section 2.1 above" in p.text for p in out.paragraphs)
    assert any("field code(s)" in n for n in doc.notes)


def test_content_control_paragraph_survives(tmp_path, profile):
    src = docx.Document()
    src.add_paragraph("Regular paragraph.")
    body = src.element.body
    sdt = parse_xml(
        f"<w:sdt {W}><w:sdtPr/><w:sdtContent>"
        "<w:p><w:r><w:t>Inside a content control.</w:t></w:r></w:p>"
        "</w:sdtContent></w:sdt>"
    )
    body.insert(list(body).index(src.paragraphs[-1]._p) + 1, sdt)
    src_path = tmp_path / "src.docx"
    src.save(src_path)

    doc, out, _ = _pipeline(src_path, profile, tmp_path)
    assert any("Inside a content control." == p.text for p in out.paragraphs)


def test_native_numpr_list_detected(tmp_path, profile):
    src = docx.Document()
    for text, lvl in (("First item", 0), ("Nested item", 1)):
        p = src.add_paragraph(text, style="List Paragraph")
        p._p.get_or_add_pPr().append(parse_xml(
            f'<w:numPr {W}><w:ilvl w:val="{lvl}"/><w:numId w:val="1"/></w:numPr>'
        ))
    src_path = tmp_path / "src.docx"
    src.save(src_path)

    doc = classify(ingest(src_path))
    items = [b for b in doc.blocks if b.label is BlockType.LIST_ITEM]
    assert len(items) == 2
    assert all(b.confidence >= CONFIDENCE_THRESHOLD for b in items)
    assert items[1].hints.list_level == 1


def test_hyperlink_url_carried(tmp_path, profile):
    src = docx.Document()
    p = src.add_paragraph("See the ")
    rid = src.part.relate_to("https://example.com/spec", RT_HYPERLINK, is_external=True)
    p._p.append(parse_xml(
        f'<w:hyperlink {W} {R} r:id="{rid}"><w:r><w:t>specification</w:t></w:r></w:hyperlink>'
    ))
    p.add_run(" for details.")
    src_path = tmp_path / "src.docx"
    src.save(src_path)

    doc, out, styled = _pipeline(src_path, profile, tmp_path)
    para = next(p for p in out.paragraphs if "specification" in p.text)
    h = para._p.find(qn("w:hyperlink"))
    assert h is not None
    out_rid = h.get(qn("r:id"))
    rel = out.part.rels[out_rid]
    assert rel.is_external and rel.target_ref == "https://example.com/spec"


def test_comment_surfaced_in_qa(tmp_path, profile):
    src = docx.Document()
    p = src.add_paragraph("Some text under discussion.")
    p._p.append(parse_xml(f'<w:r {W}><w:commentReference w:id="1"/></w:r>'))
    src_path = tmp_path / "src.docx"
    src.save(src_path)
    _inject_part(
        src_path,
        "comments.xml",
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:comments {W}><w:comment w:id="1" w:author="Reviewer A">'
        "<w:p><w:r><w:t>Please verify this figure.</w:t></w:r></w:p>"
        "</w:comment></w:comments>",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments",
    )

    doc = classify(ingest(src_path))
    assert any("Reviewer A" in n and "Please verify" in n for n in doc.notes)


def test_endnote_carried_as_footnote(tmp_path, profile):
    src = docx.Document()
    p = src.add_paragraph("Statement needing a note.")
    p._p.append(parse_xml(f'<w:r {W}><w:endnoteReference w:id="1"/></w:r>'))
    src_path = tmp_path / "src.docx"
    src.save(src_path)
    _inject_part(
        src_path,
        "endnotes.xml",
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:endnotes {W}><w:endnote w:id="1">'
        "<w:p><w:r><w:t>The endnote body text.</w:t></w:r></w:p>"
        "</w:endnote></w:endnotes>",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/endnotes",
    )

    doc, out, styled = _pipeline(src_path, profile, tmp_path)
    z = zipfile.ZipFile(styled)
    assert b"The endnote body text." in z.read("word/footnotes.xml")
    assert any("endnote(s)" in n for n in doc.notes)


def test_table_hyperlink_rewritten_not_leaked(tmp_path, profile):
    src = docx.Document()
    table = src.add_table(rows=1, cols=1)
    cell_p = table.rows[0].cells[0].paragraphs[0]
    rid = src.part.relate_to("https://example.com/in-table", RT_HYPERLINK, is_external=True)
    cell_p._p.append(parse_xml(
        f'<w:hyperlink {W} {R} r:id="{rid}"><w:r><w:t>table link</w:t></w:r></w:hyperlink>'
    ))
    src.add_paragraph("After the table.")
    src_path = tmp_path / "src.docx"
    src.save(src_path)

    doc, out, _ = _pipeline(src_path, profile, tmp_path)
    tbl = out.tables[0]._element
    # Every relationship id remaining in the carried table must resolve in the
    # OUTPUT package (the source ids are meaningless there).
    r_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    for el in tbl.iter():
        for attr, value in el.attrib.items():
            if attr.startswith(r_ns):
                assert value in out.part.rels, f"dangling rel {value} on {el.tag}"
    h = next(iter(tbl.iter(qn("w:hyperlink"))))
    assert out.part.rels[h.get(qn("r:id"))].target_ref == "https://example.com/in-table"


def test_numbered_heading_vs_list_discrimination():
    def label_of(texts, target, bold_targets=()):
        doc = Document(blocks=[
            Block(text=t) for t in texts
        ])
        for b in doc.blocks:
            if b.text in bold_targets:
                b.hints.bold = True
        classify(doc)
        return next(b for b in doc.blocks if b.text == target)

    # A run of "N." items stays an (ordered) list at high confidence.
    seq = ["1. Wear gloves", "2. Keep the area dry", "3. Check torque"]
    item = label_of(seq, "2. Keep the area dry")
    assert item.label is BlockType.LIST_NUMBER and item.confidence >= 0.85

    # A lone bold "1. Introduction" is a heading.
    b = label_of(["1. Introduction"], "1. Introduction", bold_targets=["1. Introduction"])
    assert b.label is BlockType.HEADING1

    # A lone plain "1. Introduction" is ambiguous -> flagged for QA, not silent.
    b = label_of(["1. Introduction"], "1. Introduction")
    assert b.confidence < CONFIDENCE_THRESHOLD

    # Bare "5 people attended the meeting" is body, never an H1.
    b = label_of(["5 people attended the meeting"], "5 people attended the meeting")
    assert b.label is BlockType.BODY

    # "A. Smith et al. (2020) argue" is prose, not a list item (ordered or not).
    b = label_of(["A. Smith et al. (2020) argue"], "A. Smith et al. (2020) argue")
    assert b.label not in (BlockType.LIST_ITEM, BlockType.LIST_NUMBER)
