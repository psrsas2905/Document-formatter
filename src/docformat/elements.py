"""Assemble template scaffolding: headers/footers, page numbers, captions, TOC.

Operates on the already-styled .docx produced by apply_styles:
  - page size/margins from the profile
  - header/footer text with {page}/{pages}/{doc_title}/{version}/{classification}
    tokens — {page}/{pages} become real PAGE/NUMPAGES fields (lxml OOXML)
  - the template's own header/footer content (brand logo etc.) is PRESERVED;
    profile text is appended alongside it
  - TOC and List of Figures inserted as dirty field codes; LibreOffice (export
    stage) or Word refreshes them into real entries

Field codes need raw OOXML that python-docx has no API for, hence lxml here.
"""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm

PAGE_SIZES_MM = {"A4": (210, 297), "Letter": (215.9, 279.4)}
FIELD_TOKENS = {"{page}": "PAGE", "{pages}": "NUMPAGES"}


def add_elements(docx_path: str | Path, profile, model=None) -> Path:
    """Add headers/footers/page numbers/TOC to an existing styled .docx.

    `model` is the pipeline's Document; when given, leftover template
    placeholders (like "[Client Name]") are appended to its QA notes.
    """
    docx_path = Path(docx_path)
    doc = docx.Document(str(docx_path))
    raw = profile.raw

    tokens = {
        "{doc_title}": doc.core_properties.title or "",
        "{version}": str(raw.get("output", {}).get("version", "")),
        "{classification}": str(raw.get("footer", {}).get("classification", "")),
        "{date}": date.today().strftime("%d %B %Y"),
    }

    _apply_page_setup(doc, raw.get("page", {}))
    _write_header_footer(doc, raw, tokens)
    _auto_number_captions(doc, profile)
    _insert_front_matter(doc, raw)
    if model is not None:
        _note_unfilled_placeholders(doc, model)

    doc.save(docx_path)
    return docx_path


# --- page setup -------------------------------------------------------------


def _apply_page_setup(doc, page: dict) -> None:
    size = PAGE_SIZES_MM.get(page.get("size", ""), None)
    margins = page.get("margins_mm", {})
    for section in doc.sections:
        if size:
            section.page_width, section.page_height = Mm(size[0]), Mm(size[1])
        for side in ("top", "bottom", "left", "right"):
            if side in margins:
                setattr(section, f"{side}_margin", Mm(margins[side]))


# --- header / footer --------------------------------------------------------


def _write_header_footer(doc, raw: dict, tokens: dict) -> None:
    from .apply import _paragraph_styles

    header_spec = raw.get("header", {})
    footer_spec = raw.get("footer", {})
    styles = _paragraph_styles(doc)

    # Placeholder substitution inside the template's OWN content — headers,
    # footers and the (kept) cover page — e.g. "[Document Title]" -> the real
    # title. Configured under a top-level `replace:` (header.replace /
    # footer.replace also accepted).
    replacements = {
        **raw.get("replace", {}),
        **header_spec.get("replace", {}),
        **footer_spec.get("replace", {}),
    }
    if replacements:
        from docx.text.paragraph import Paragraph

        parts = [doc.element.body]  # cover page content lives in the body
        for section in doc.sections:
            for part in (section.header, section.footer,
                         section.first_page_header, section.first_page_footer):
                parts.append(part._element)
        for root in parts:
            # Walk every paragraph, including those nested in tables/text
            # boxes — brand covers and headers are often laid out that way.
            for p_el in root.iter(qn("w:p")):
                _replace_placeholders(Paragraph(p_el, doc), replacements, tokens)

    for section in doc.sections:
        if header_spec.get("text"):
            _append_marginal_text(
                section.header, header_spec["text"], tokens, style=styles.get("Header"),
                align=WD_ALIGN_PARAGRAPH.RIGHT,
            )
        if footer_spec.get("text"):
            _append_marginal_text(
                section.footer, footer_spec["text"], tokens, style=styles.get("Footer"),
                align=WD_ALIGN_PARAGRAPH.CENTER,
            )
        if not header_spec.get("show_on_first_page", True):
            section.different_first_page_header_footer = True
            # First-page header/footer stay empty (created blank on demand).
            _ = section.first_page_header, section.first_page_footer

    doc.settings.element  # ensure settings part exists before saving


def _replace_placeholders(para, replacements: dict, tokens: dict) -> None:
    """Swap template placeholders (e.g. '[Document Title]') for token-expanded
    values inside existing header/footer/cover runs.

    Works on the run concatenation (NOT para.text, which also includes
    hyperlink-nested runs) and rewrites only the runs the placeholder spans,
    so surrounding runs keep their own formatting and hyperlink text is never
    duplicated."""
    runs = para.runs
    if not runs:
        return
    for placeholder, value in replacements.items():
        for token, actual in tokens.items():
            value = value.replace(token, actual)
        while True:
            texts = [r.text for r in runs]
            full = "".join(texts)
            start = full.find(placeholder)
            if start == -1:
                break
            end = start + len(placeholder)
            spans = []  # (run index, slice start, slice end) touched by the match
            pos = 0
            for i, t in enumerate(texts):
                if pos + len(t) > start and pos < end:
                    spans.append((i, max(0, start - pos), min(len(t), end - pos)))
                pos += len(t)
            first_i, f_start, f_end = spans[0]
            if len(spans) == 1:
                runs[first_i].text = (
                    texts[first_i][:f_start] + value + texts[first_i][f_end:]
                )
            else:
                runs[first_i].text = texts[first_i][:f_start] + value
                for i, _, _ in spans[1:-1]:
                    runs[i].text = ""
                last_i, _, l_end = spans[-1]
                runs[last_i].text = texts[last_i][l_end:]


def _append_marginal_text(container, text: str, tokens: dict, style, align) -> None:
    """Append one paragraph of profile text to a header/footer, preserving
    whatever the template already put there (e.g. the brand logo)."""
    container.is_linked_to_previous = False
    para = container.add_paragraph()
    if style is not None:
        para.style = style
    para.alignment = align

    for token in tokens:
        text = text.replace(token, tokens[token])
    # Split on {page}/{pages} so those become live fields, not literals.
    for piece in re.split(r"(\{pages?\})", text):
        if piece in FIELD_TOKENS:
            _add_field_run(para, FIELD_TOKENS[piece])
        elif piece:
            para.add_run(piece)


def _add_field_run(para, instr: str, placeholder: str = "1") -> None:
    """Append a simple field (PAGE / NUMPAGES / SEQ ...) as raw OOXML."""
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), f" {instr} ")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = placeholder  # shown until fields refresh
    run.append(t)
    fld.append(run)
    para._p.append(fld)


# --- captions ----------------------------------------------------------------

CAPTION_NUM_RE = re.compile(r"^(Figure|Table)\s+\d+(\s*[:.\-–]?\s*)(.*)$", re.IGNORECASE)


def _auto_number_captions(doc, profile) -> None:
    """Replace static 'Figure 1'/'Table 2' numbers in caption paragraphs with
    SEQ fields, so numbering self-heals and the List of Figures can find them."""
    caption_style = profile.style_map.get("Caption", "Caption")
    counters: dict[str, int] = {}
    for para in doc.paragraphs:
        if para.style is None or para.style.name != caption_style:
            continue
        m = CAPTION_NUM_RE.match(para.text)
        if not m:
            continue
        kind = m.group(1).capitalize()
        counters[kind] = counters.get(kind, 0) + 1
        sep, rest = m.group(2) or " ", m.group(3)

        # Replace only the TEXT runs; a caption paragraph may also carry
        # non-text runs (an inline image) that must survive in place.
        p_el = para._p
        text_runs = [r for r in para.runs if r._element.find(qn("w:t")) is not None]
        if not text_runs:
            continue
        insert_at = list(p_el).index(text_runs[0]._element)
        for run in text_runs:
            p_el.remove(run._element)

        created = [para.add_run(f"{kind} ")._element]
        _add_field_run(para, rf"SEQ {kind} \* ARABIC", placeholder=str(counters[kind]))
        created.append(p_el[-1])
        created.append(para.add_run(f"{sep}{rest}" if rest else "")._element)
        for el in reversed(created):
            p_el.remove(el)
            p_el.insert(insert_at, el)


# --- TOC / List of Figures --------------------------------------------------


def _insert_front_matter(doc, raw: dict) -> None:
    """Insert TOC and List of Figures after the document title (or at start)."""
    h1_style = raw.get("style_map", {}).get("Heading1", "Heading 1")
    anchor = _front_matter_anchor(doc, {h1_style, "Heading 1"})

    toc = raw.get("toc", {})
    lof = raw.get("list_of_figures", {})

    if toc.get("enabled"):
        max_level = int(toc.get("max_level", 3))
        anchor = _insert_field_block(
            doc, anchor, toc.get("title", "Table of Contents"),
            rf'TOC \o "1-{max_level}" \h \z \u',
        )
    if lof.get("enabled"):
        anchor = _insert_field_block(
            doc, anchor, lof.get("title", "List of Figures"),
            r'TOC \h \z \c "Figure"',
        )


PLACEHOLDER_RE = re.compile(r"\[[^\[\]\n]{1,60}\]")


def _note_unfilled_placeholders(doc, model) -> None:
    """QA-note any template placeholders like '[Client Name]' still unfilled
    (in the kept cover page, headers or footers) so a human fills them in."""
    from docx.text.paragraph import Paragraph

    roots = [doc.element.body] + [
        part._element
        for section in doc.sections
        for part in (section.header, section.footer,
                     section.first_page_header, section.first_page_footer)
    ]
    leftover: dict[str, int] = {}
    for root in roots:
        for p_el in root.iter(qn("w:p")):
            for m in PLACEHOLDER_RE.findall(Paragraph(p_el, doc).text):
                leftover[m] = leftover.get(m, 0) + 1
    for placeholder, count in sorted(leftover.items()):
        model.notes.append(
            f"Template placeholder {placeholder} is still unfilled "
            f"({count} occurrence(s)) — fill it in, or map it under `replace:` "
            "in the profile."
        )


def _front_matter_anchor(doc, h1_names: set[str]):
    """Element after which front matter goes: the poured document's H1 title
    (whatever style name the profile maps it to).

    With a kept cover page but no H1 in the content, fall back to the cover's
    section-break paragraph so the TOC never lands ABOVE the cover.
    """
    for para in doc.paragraphs:
        if para.style is not None and para.style.name in h1_names:
            return para._p
    for para in doc.paragraphs:  # cover section break, if a cover was kept
        pPr = para._p.find(qn("w:pPr"))
        if pPr is not None and pPr.find(qn("w:sectPr")) is not None:
            return para._p
    return None


def _insert_field_block(doc, anchor, title: str, instr: str):
    """Insert '<title paragraph><dirty field paragraph>' after anchor.

    Returns the field paragraph element (the next block's anchor).
    """
    body = doc.element.body

    title_p = _make_paragraph(doc, title, _pick_style(doc, "TOC Heading"))
    field_p = OxmlElement("w:p")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), f" {instr} ")
    fld.set(qn("w:dirty"), "true")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "(updated on export)"
    run.append(t)
    fld.append(run)
    field_p.append(fld)

    if anchor is not None:
        anchor.addnext(title_p)
    else:
        body.insert(0, title_p)
    title_p.addnext(field_p)
    return field_p


def _make_paragraph(doc, text: str, style):
    para = doc.add_paragraph(text)
    if style is not None:
        para.style = style  # style OBJECT — survives nonstandard internal names
    para._p.getparent().remove(para._p)
    return para._p


def _pick_style(doc, preferred: str, fallback: str | None = None):
    """Style object for the first defined name, or None (plain paragraph).

    Falling back to a heading style would put TOC/LOF titles inside the TOC
    itself, so templates without the preferred style get an unstyled title.
    """
    from .apply import _paragraph_styles

    styles = _paragraph_styles(doc)
    style = styles.get(preferred)
    if style is None and fallback is not None:
        style = styles.get(fallback)
    return style
