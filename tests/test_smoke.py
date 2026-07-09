"""Smoke tests — expand into real regression tests against samples/ as stages land."""

from docformat.template import load_profile


def test_profile_loads():
    profile = load_profile("config/template_profile.example.yaml")
    assert profile.style_map["Heading1"] == "Heading 1"
    assert profile.raw["page"]["size"] in {"A4", "Letter"}


def test_ingest_sample():
    from docformat.ingest import ingest

    doc = ingest("samples/input_messy.docx")
    assert len(doc.blocks) == 21  # empty spacing paragraphs dropped
    texts = [b.text for b in doc.blocks]

    title = doc.blocks[texts.index("ACME WIDGET INSTALLATION GUIDE")]
    assert title.hints.font_size_pt == 18.0 and title.hints.bold and title.hints.all_caps

    bullets = [b for b in doc.blocks if b.text.startswith(("•", "-", "*"))]
    assert len(bullets) == 3 and all(b.hints.is_list_marker for b in bullets)

    sub = doc.blocks[texts.index("a) Torque each bolt to the value shown below.")]
    assert sub.hints.is_list_marker and sub.hints.list_level == 1

    styled = doc.blocks[texts.index("Maintenance Schedule")]
    assert styled.hints.existing_style == "Heading 2"


def test_classify_sample():
    from docformat.classify import CONFIDENCE_THRESHOLD, classify
    from docformat.ingest import ingest
    from docformat.models import BlockType

    doc = classify(ingest("samples/input_messy.docx"))
    by_text = {b.text: b for b in doc.blocks}

    assert doc.title == "ACME WIDGET INSTALLATION GUIDE"
    assert by_text["ACME WIDGET INSTALLATION GUIDE"].label is BlockType.HEADING1
    assert by_text["Overview"].label is BlockType.HEADING2
    assert by_text["Figure 1: Wiring diagram for single-phase supply"].label is BlockType.CAPTION
    assert by_text["Table 2 - Torque settings by bolt size"].label is BlockType.CAPTION
    assert by_text["• Wear insulated gloves rated to 1000 V"].label is BlockType.LIST_ITEM
    assert by_text["3. Connect the signal cable to port A."].label is BlockType.LIST_ITEM

    trusted = by_text["Maintenance Schedule"]
    assert trusted.label is BlockType.HEADING2 and trusted.confidence == 1.0

    ambiguous = by_text["Safety Precautions"]
    assert ambiguous.label in {BlockType.HEADING2, BlockType.HEADING3}
    assert ambiguous.confidence < CONFIDENCE_THRESHOLD  # must reach the QA report

    assert all(b.label is not BlockType.UNKNOWN for b in doc.blocks)


def test_apply_styles(tmp_path):
    import docx

    from docformat.apply import apply_styles
    from docformat.classify import classify
    from docformat.ingest import ingest
    from docformat.template import load_profile

    profile = load_profile("config/template_profile.example.yaml")
    doc = classify(ingest("samples/input_messy.docx"))
    out = apply_styles(doc, profile, tmp_path / "styled.docx")

    result = docx.Document(str(out))
    styles = [p.style.name for p in result.paragraphs]
    assert styles[0] == "Heading 1"
    assert "List Bullet" in styles and "Caption" in styles and "Quote" in styles
    assert "List Bullet 2" in styles  # nested a)/b) items pick the level variant

    for para in result.paragraphs:
        assert "\t" not in para.text
        for run in para.runs:  # no direct formatting survives
            assert run.font.size is None and run.font.bold is None and run.font.italic is None
    markers = ("•", "-", "*", "1.", "a)")
    assert not any(p.text.startswith(markers) for p in result.paragraphs)


def _run_pipeline(tmp_path):
    from docformat.apply import apply_styles
    from docformat.classify import classify
    from docformat.elements import add_elements
    from docformat.ingest import ingest
    from docformat.template import load_profile

    profile = load_profile("config/template_profile.example.yaml")
    doc = classify(ingest("samples/input_messy.docx"))
    styled = apply_styles(doc, profile, tmp_path / "input_messy_formatted.docx")
    add_elements(styled, profile)
    return doc, styled


def test_golden_regression(tmp_path):
    import docx

    _, styled = _run_pipeline(tmp_path)
    result = docx.Document(str(styled))
    got = [f"{p.style.name}\t{p.text}" for p in result.paragraphs]
    expected = open("tests/golden/input_messy_styles.tsv", encoding="utf-8").read().splitlines()
    assert got == expected


def test_brand_template_preserved(tmp_path):
    import docx

    _, styled = _run_pipeline(tmp_path)
    result = docx.Document(str(styled))
    header = result.sections[0].header
    blips = header._element.findall(
        ".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
    )
    assert len(blips) == 1  # the org logo survived into the output header
    assert any("ACME CORP" in p.text for p in header.paragraphs)


def test_qa_report(tmp_path):
    from docformat.qa import write_report

    doc, _ = _run_pipeline(tmp_path)
    report = write_report(doc, tmp_path / "qa_report.md")
    text = report.read_text(encoding="utf-8")
    assert "Safety Precautions" in text  # the one deliberately ambiguous block
    assert "No level jumps detected" in text


def test_export_pdf(tmp_path):
    import subprocess

    import pytest

    from docformat.export import export_pdf, soffice_available

    if not soffice_available():
        pytest.skip("LibreOffice not installed")
    _, styled = _run_pipeline(tmp_path)
    pdf = export_pdf(styled, tmp_path)
    assert pdf.exists() and pdf.stat().st_size > 10_000
    data = pdf.read_bytes()
    assert b"/Marked true" in data  # tagged PDF
    text = subprocess.run(
        ["pdftotext", str(pdf), "-"], capture_output=True, text=True
    ).stdout
    assert "Table of Contents" in text
    assert "Overview" in text.split("Table of Contents")[1][:600]  # TOC populated


def test_apply_field_values():
    from docformat.template import apply_field_values, load_profile

    profile = load_profile("config/template_profile.redlotus.yaml")
    apply_field_values(profile, {"Client Name": "Acme", "[Custom]": "x"})
    assert profile.raw["replace"]["[Client Name]"] == "Acme"
    assert profile.raw["replace"]["[Custom]"] == "x"
    assert profile.raw["replace"]["[Document Title]"] == "{doc_title}"  # untouched


def test_numbered_heading_heuristic():
    from docformat.classify import classify
    from docformat.ingest import ingest
    from docformat.models import BlockType

    doc = classify(ingest("samples/input_rich.docx"))
    by_text = {b.text: b for b in doc.blocks}
    assert by_text["1 Introduction"].label is BlockType.HEADING1
    assert by_text["1.1 Governing Equation"].label is BlockType.HEADING2
    assert by_text["2 Hardware Layout"].label is BlockType.HEADING1
    assert all(
        b.confidence >= 0.85
        for b in doc.blocks
        if b.label in {BlockType.HEADING1, BlockType.HEADING2}
    )


def test_context_heading_heuristic():
    from docformat.classify import CONFIDENCE_THRESHOLD, classify
    from docformat.models import Block, BlockType, Document

    long = "x" * 150
    doc = Document(blocks=[Block(text=long), Block(text="Lost Heading"), Block(text=long)])
    classify(doc)
    middle = doc.blocks[1]
    assert middle.label is BlockType.HEADING2
    assert middle.confidence < CONFIDENCE_THRESHOLD  # level is a guess -> QA


def test_convert_txt_input(tmp_path):
    import docx

    from docformat.convert import ensure_docx

    src = tmp_path / "draft.txt"
    src.write_text("Title Line\n\nBody paragraph one\ncontinued here.\n\nSecond para.\n")
    out = ensure_docx(src, tmp_path / "conv")
    texts = [p.text for p in docx.Document(str(out)).paragraphs]
    assert texts == ["Title Line", "Body paragraph one continued here.", "Second para."]


def test_convert_doc_input(tmp_path):
    import pytest

    from docformat.convert import ensure_docx
    from docformat.export import soffice_available
    from docformat.ingest import ingest

    if not soffice_available():
        pytest.skip("LibreOffice not installed")
    import subprocess

    subprocess.run(
        ["soffice", "--headless", "--convert-to", "doc", "--outdir", str(tmp_path),
         "samples/input_messy.docx"],
        check=True, capture_output=True,
    )
    out = ensure_docx(tmp_path / "input_messy.doc", tmp_path / "conv")
    assert out.suffix == ".docx" and len(ingest(out).blocks) > 15


def test_convert_rejects_unknown(tmp_path):
    import pytest

    from docformat.convert import ensure_docx

    bad = tmp_path / "x.xyz"
    bad.write_text("data")
    with pytest.raises(ValueError, match="Unsupported input type"):
        ensure_docx(bad, tmp_path)


def test_preset_profiles_load():
    from docformat.template import load_profile

    apa = load_profile("config/template_profile.apa.yaml")
    assert apa.raw["page"]["size"] == "Letter" and not apa.raw["toc"]["enabled"]
    gbt = load_profile("config/template_profile.gbt7713.yaml")
    assert gbt.raw["page"]["margins_mm"]["left"] == 31.7 and gbt.raw["toc"]["enabled"]
    for profile in (apa, gbt):
        assert set(profile.style_map) == {
            "Heading1", "Heading2", "Heading3", "Body", "Caption", "ListItem", "Quote",
        }
