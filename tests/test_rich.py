"""Content-preservation regression: equations, images, tables, footnotes,
inline emphasis must survive the pipeline (samples/input_rich.docx)."""

import zipfile

import docx
import pytest

from docformat.apply import apply_styles
from docformat.classify import classify
from docformat.elements import add_elements
from docformat.ingest import ingest
from docformat.models import BlockType
from docformat.template import load_profile


@pytest.fixture(scope="module")
def result(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("rich")
    profile = load_profile("config/template_profile.example.yaml")
    doc = classify(ingest("samples/input_rich.docx"))
    styled = apply_styles(doc, profile, tmp / "out.docx")
    add_elements(styled, profile, doc)
    return doc, styled


def test_equation_carried_verbatim(result):
    _, styled = result
    xml = zipfile.ZipFile(styled).read("word/document.xml").decode()
    assert xml.count("<m:oMath") == 1
    assert "<m:rad>" in xml  # the square root survived structurally


def test_images_reembedded_with_size(result):
    _, styled = result
    out = docx.Document(str(styled))
    shapes = out.inline_shapes
    assert len(shapes) == 2
    widths = sorted(s.width.emu for s in shapes)
    assert widths == [139700, 914400]  # inline 11pt and figure 72pt, as authored


def test_table_carried_and_restyled(result):
    _, styled = result
    out = docx.Document(str(styled))
    assert len(out.tables) == 1
    table = out.tables[0]
    assert table.style.name == "Table Grid"  # profile's table_style
    assert table.rows[0].cells[0].text == "Parameter"
    assert table.rows[2].cells[2].text == "12.5 rad/s"


def test_footnote_carried(result):
    doc, styled = result
    z = zipfile.ZipFile(styled)
    assert "word/footnotes.xml" in z.namelist()
    assert b"Sampling above 2x" in z.read("word/footnotes.xml")
    assert b"footnoteReference" in z.read("word/document.xml")
    assert any("footnote" in n for n in doc.notes)


def test_inline_emphasis_preserved(result):
    _, styled = result
    out = docx.Document(str(styled))
    para = next(p for p in out.paragraphs if p.text.startswith("This note"))
    flags = {r.text: (bool(r.font.bold), bool(r.font.italic)) for r in para.runs if r.text}
    assert flags["critically damped"] == (True, False)
    assert flags["only valid"] == (False, True)
    assert flags["This note describes the "] == (False, False)


def test_table_block_classified(result):
    doc, _ = result
    tables = [b for b in doc.blocks if b.label is BlockType.TABLE]
    assert len(tables) == 1 and tables[0].confidence == 1.0


def test_character_formatting_preserved(result):
    """Sub/superscript (meaning) plus underline/strike (inline emphasis) survive."""
    _, styled = result
    out = docx.Document(str(styled))
    para = next(p for p in out.paragraphs if p.text.startswith("The coolant is"))
    runs = {r.text: r for r in para.runs if r.text}

    # H₂O and v² — the digit runs carry vertAlign, not the letters.
    assert runs["H"].font.subscript in (None, False)
    subscripts = [r.text for r in para.runs if r.font.subscript]
    superscripts = [r.text for r in para.runs if r.font.superscript]
    assert subscripts == ["2"] and superscripts == ["2"]

    assert bool(runs["thermal budget"].font.underline) is True
    assert bool(runs["old 5 W limit"].font.strike) is True
    # Highlighter marks (author annotations) survive.
    from docx.enum.text import WD_COLOR_INDEX
    assert runs["Confirm the tolerance"].font.highlight_color == WD_COLOR_INDEX.YELLOW
    # Plain text stays plain.
    assert not runs["The coolant is "].font.underline
    assert not runs["The coolant is "].font.strike
    assert runs["The coolant is "].font.highlight_color is None
