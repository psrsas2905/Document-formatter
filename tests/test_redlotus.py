"""Real-world template trial: RedLotus master template.

This template exercises the messy cases sanitized demo templates don't:
dangling style references, duplicate style definitions, nonstandard internal
style names, no Body/Caption/Quote styles, branding inside a header table,
and a [Document Title] placeholder.
"""

from pathlib import Path

import docx
import pytest

from docformat.apply import apply_styles
from docformat.classify import classify
from docformat.elements import add_elements
from docformat.ingest import ingest
from docformat.models import BlockType
from docformat.qa import write_report
from docformat.template import load_profile

TEMPLATE = Path("templates/RedLotus_Master_Template.docx")
A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


@pytest.fixture(scope="module")
def result(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("redlotus")
    profile = load_profile("config/template_profile.redlotus.yaml")
    doc = classify(ingest("samples/input_messy.docx"))
    styled = apply_styles(doc, profile, tmp / "out.docx")
    add_elements(styled, profile)
    qa = write_report(doc, tmp / "qa_report.md")
    return doc, docx.Document(str(styled)), qa


def test_ingests_template_with_dangling_styles():
    # The template itself references undefined style ids; ingest must not crash.
    doc = classify(ingest(str(TEMPLATE)))
    assert len(doc.blocks) > 40
    assert all(b.label is not BlockType.UNKNOWN for b in doc.blocks)


def test_brand_header_and_footer_survive(result):
    _, out, _ = result
    header = out.sections[0].header
    assert len(header._element.findall(f".//{A_NS}blip")) == 2  # logo banner
    footer_text = "".join(p.text for p in out.sections[0].footer.paragraphs)
    assert "RedLotus Pharmtech" in footer_text


def test_header_placeholder_replaced(result):
    _, out, _ = result
    header_xml = out.sections[0].header._element.xml
    assert "[Document Title]" not in header_xml
    assert "ACME WIDGET INSTALLATION GUIDE" in header_xml


def test_defined_styles_applied_and_missing_ones_reported(result):
    doc, out, qa = result
    used = {p.style.name for p in out.paragraphs if p.style is not None}
    assert {"Heading 1", "Heading 2", "List Paragraph", "Strong"} <= used
    # 'Normal' isn't defined in this template: paragraphs fall back to default
    # formatting and the mismatch must reach the QA report.
    assert any("'Normal'" in n for n in doc.notes)
    assert "Template / profile warnings" in qa.read_text(encoding="utf-8")
